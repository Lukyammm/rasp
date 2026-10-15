import json
import os
import sys
import tempfile
import time
import unittest
from contextlib import contextmanager
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("RASP_DOWNLOADS_DIR", tempfile.mkdtemp())
os.environ.setdefault("RASP_LOGS_DIR", tempfile.mkdtemp())

from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402

import database  # noqa: E402
from scrapers import comments as cm  # noqa: E402


def bucket(limit=50, seconds=60):
    return cm.CommentBucket(limit=limit, deadline=time.monotonic() + seconds)


class FakePage:
    def __init__(self, closed=False):
        self.closed = closed

    def is_closed(self):
        return self.closed

    def wait_for_timeout(self, ms):
        pass


class BucketTests(unittest.TestCase):
    def test_dedupe_limit_and_stable_ids(self):
        b = bucket(limit=3)
        self.assertTrue(b.add("1", "@a", "oi"))
        self.assertFalse(b.add("1", "@a", "oi de novo"))  # mesmo id
        self.assertTrue(b.add("", "@b", "sem   id"))
        self.assertFalse(b.add("", "@b", "sem id"))  # mesmo conteúdo -> mesmo hash
        self.assertFalse(b.add("9", "", ""))  # vazio
        self.assertTrue(b.add("2", "@c", "tres"))
        self.assertTrue(b.full)
        self.assertFalse(b.add("3", "@d", "excede"))
        self.assertEqual(b.items[1]["text"], "sem id")
        self.assertTrue(b.items[1]["id"].startswith("h"))

    def test_deadline(self):
        b = bucket(seconds=-1)
        self.assertTrue(b.should_stop())
        self.assertTrue(b.timed_out)


class ParserTests(unittest.TestCase):
    def test_youtube_entities_and_next_token(self):
        doc = {
            "frameworkUpdates": {"entityBatchUpdate": {"mutations": [
                {"payload": {"commentEntityPayload": {
                    "properties": {"commentId": "c1", "content": {"content": "ótimo"}, "publishedTime": "há 2 dias", "replyLevel": 0},
                    "author": {"displayName": "@fulano"}, "toolbar": {"likeCountNotliked": "1,2 mil", "replyCount": "3"}}}},
                {"payload": {"commentEntityPayload": {
                    "properties": {"commentId": "r1", "content": {"content": "resposta"}, "replyLevel": 1},
                    "author": {"displayName": "@x"}, "toolbar": {}}}},
            ]}},
            "onResponseReceivedEndpoints": [{"appendContinuationItemsAction": {"continuationItems": [
                {"commentThreadRenderer": {"replies": {"continuationItemRenderer": {"continuationCommand": {"token": "REPLIES"}}}}},
                {"continuationItemRenderer": {"continuationEndpoint": {"continuationCommand": {"token": "NEXT"}}}},
            ]}}],
        }
        b = bucket()
        self.assertEqual(cm._yt_collect(doc, b), 1)
        self.assertEqual(b.items[0]["likes"], 1200)
        self.assertEqual(b.items[0]["replies"], 3)
        self.assertEqual(b.items[0]["date"], "há 2 dias")
        self.assertEqual(cm._yt_next_token(doc), "NEXT")
        self.assertEqual(cm._yt_next_token({"onResponseReceivedEndpoints": []}), "")

    def test_youtube_video_id(self):
        self.assertEqual(cm._yt_video_id("https://www.youtube.com/watch?v=dQw4w9WgXcQ&t=1"), "dQw4w9WgXcQ")
        self.assertEqual(cm._yt_video_id("https://www.youtube.com/shorts/dQw4w9WgXcQ"), "dQw4w9WgXcQ")
        self.assertEqual(cm._yt_video_id("https://www.youtube.com/@canal"), "")

    def test_x_replies_only_from_conversation(self):
        def tweet(tid, conv, reply_to, text, handle="u"):
            return {"itemContent": {"tweet_results": {"result": {"rest_id": tid, "core": {"user_results": {"result": {"core": {"screen_name": handle}}}},
                    "legacy": {"id_str": tid, "conversation_id_str": conv, "in_reply_to_status_id_str": reply_to,
                               "full_text": text, "favorite_count": 2, "reply_count": 0, "created_at": "Wed Sep 09 22:38:06 +0000 2026"}}}}}
        body = json.dumps({"data": {"entries": [
            tweet("100", "100", None, "post original"),
            tweet("101", "100", "100", "@dono @outro concordo &amp; muito"),
            tweet("102", "100", "101", "@u tréplica"),
            tweet("200", "999", None, "recomendação de outro assunto"),
        ]}})
        b = bucket()
        self.assertEqual(cm._x_collect(body, "100", b), 2)
        by_id = {c["id"]: c for c in b.items}
        self.assertEqual(by_id["101"]["text"], "concordo & muito")
        self.assertEqual(by_id["101"]["parent_id"], "")
        self.assertEqual(by_id["102"]["parent_id"], "101")

    def test_x_replies_when_focal_post_is_itself_a_reply(self):
        # Quando o post focal é resposta, seu conversation_id_str aponta para a raiz da thread (ex: 50) e não para o focal_id (100)
        def tweet(tid, conv, reply_to, text, handle="u"):
            return {"itemContent": {"tweet_results": {"result": {"rest_id": tid, "core": {"user_results": {"result": {"core": {"screen_name": handle}}}},
                    "legacy": {"id_str": tid, "conversation_id_str": conv, "in_reply_to_status_id_str": reply_to,
                               "full_text": text, "favorite_count": 5, "reply_count": 1, "created_at": "Wed Sep 09 22:38:06 +0000 2026"}}}}}
        # 50 é o tweet raiz da conversa
        # 100 é o post focal (resposta ao 50, portanto conv="50")
        # 101 é resposta direta ao 100 (in_reply_to="100", conv="50")
        # 102 é tréplica à resposta 101 (in_reply_to="101", conv="50")
        # 99 é resposta ao 50 (ramo irmão, não deve ser coletada)
        # 200 é tweet não relacionado
        body = json.dumps({"data": {"entries": [
            tweet("100", "50", "50", "post focal que é resposta ao 50"),
            tweet("101", "50", "100", "@focal resposta direta ao post focal"),
            tweet("102", "50", "101", "@resp réplica à resposta"),
            tweet("99", "50", "50", "@raiz outra resposta à raiz (irmão)"),
            tweet("200", "999", None, "não relacionado"),
        ]}})
        b = bucket()
        accepted_ids = {"100"}
        collected = cm._x_collect(body, "100", b, accepted_ids)
        self.assertEqual(collected, 2)
        by_id = {c["id"]: c for c in b.items}
        self.assertIn("101", by_id)
        self.assertIn("102", by_id)
        self.assertNotIn("100", by_id)
        self.assertNotIn("99", by_id)
        self.assertNotIn("200", by_id)
        self.assertEqual(by_id["101"]["parent_id"], "")
        self.assertEqual(by_id["102"]["parent_id"], "101")
        self.assertIn("101", accepted_ids)
        self.assertIn("102", accepted_ids)

    def test_x_replies_out_of_order_and_integer_in_reply_to(self):
        # Valida que respostas fora de ordem (filho antes do pai) e IDs inteiros são resolvidos
        def tweet(tid, conv, reply_to, text, reply_to_int=None):
            leg = {"id_str": tid, "conversation_id_str": conv, "full_text": text,
                   "favorite_count": 1, "reply_count": 0, "created_at": "Wed Sep 09 22:38:06 +0000 2026"}
            if reply_to is not None:
                leg["in_reply_to_status_id_str"] = reply_to
            if reply_to_int is not None:
                leg["in_reply_to_status_id"] = reply_to_int
            return {"itemContent": {"tweet_results": {"result": {"rest_id": tid, "core": {"user_results": {"result": {"core": {"screen_name": "u"}}}},
                    "legacy": leg}}}}

        # 102 (sub-resposta a 101) aparece antes de 101 (resposta direta a 100 com ID inteiro)
        body = json.dumps({"data": {"entries": [
            tweet("102", "50", "101", "sub-resposta"),
            tweet("101", "50", None, "resposta com int in_reply_to", reply_to_int=100),
        ]}})
        b = bucket()
        collected = cm._x_collect(body, "100", b)
        self.assertEqual(collected, 2)
        by_id = {c["id"]: c for c in b.items}
        self.assertEqual(by_id["101"]["parent_id"], "")
        self.assertEqual(by_id["102"]["parent_id"], "101")

        # Comentários pré-existentes no bucket também aceitam novas sub-respostas
        body2 = json.dumps({"data": {"entries": [
            tweet("103", "50", "102", "tréplica à 102"),
        ]}})
        collected2 = cm._x_collect(body2, "100", b)
        self.assertEqual(collected2, 1)
        self.assertEqual([c["id"] for c in b.items], ["101", "102", "103"])
        self.assertEqual(b.items[-1]["parent_id"], "102")

    def test_tiktok_root_comment_has_no_parent(self):
        b = bucket()
        cm._tt_collect({"comments": [{"cid": "1", "text": "a", "reply_id": "0", "user": {"unique_id": "z"}, "create_time": 1788996994},
                                     {"cid": "2", "text": "b", "reply_id": "1", "user": {"unique_id": "y"}}]}, b)
        self.assertEqual([c["parent_id"] for c in b.items], ["", "1"])
        self.assertEqual(b.items[0]["author"], "@z")


class CollectPostTests(unittest.TestCase):
    def run_with(self, fetcher, page=None, limit=50):
        with mock.patch.dict(cm.FETCHERS, {"x": fetcher}), mock.patch.object(cm, "RETRY_WAIT", (0, 0, 0)):
            return cm.collect_post_comments("x", page or FakePage(), {"id": 1, "link": "https://x.com/a/status/1"}, limit, {})

    def test_transient_error_is_retried(self):
        calls = []

        def fetcher(page, post, b, state):
            calls.append(1)
            if len(calls) < 3:
                raise cm.TransientError("HTTP 503")
            b.add("1", "@a", "ok")
            b.exhausted = True

        status, items, msg = self.run_with(fetcher)
        self.assertEqual((status, len(items), len(calls)), ("done", 1, 3))

    def test_gives_up_after_attempts_keeping_nothing(self):
        def fetcher(page, post, b, state):
            raise cm.TransientError("HTTP 500")
        status, items, msg = self.run_with(fetcher)
        self.assertEqual((status, items), ("error", []))
        self.assertIn("HTTP 500", msg)

    def test_partial_results_survive_failures(self):
        def fetcher(page, post, b, state):
            b.add(str(len(b.items)), "@a", f"c{len(b.items)}")
            raise cm.TransientError("caiu no meio")
        status, items, msg = self.run_with(fetcher)
        self.assertEqual(status, "partial")
        self.assertEqual(len(items), 3)  # uma por tentativa, sem duplicar

    def test_login_and_unavailable_are_not_retried(self):
        calls = []

        def login(page, post, b, state):
            calls.append(1)
            raise cm.LoginRequired("conecte")
        self.assertEqual(self.run_with(login)[0], "login")
        self.assertEqual(len(calls), 1)

        def gone(page, post, b, state):
            raise cm.PostUnavailable("removido")
        self.assertEqual(self.run_with(gone)[:3:2], ("error", "removido"))

    def test_disabled_and_empty(self):
        def disabled(page, post, b, state):
            b.disabled = True
        self.assertEqual(self.run_with(disabled)[0], "disabled")

        def empty(page, post, b, state):
            b.exhausted = True
        self.assertEqual(self.run_with(empty)[0], "empty")

    def test_dead_browser_raises_only_without_data(self):
        def dies(page, post, b, state):
            page.closed = True
            raise RuntimeError("Target page, context or browser has been closed")
        with self.assertRaises(cm.BrowserDead):
            self.run_with(dies)

        def dies_after_data(page, post, b, state):
            b.add("1", "@a", "salvo")
            page.closed = True
            raise RuntimeError("Target closed")
        self.assertEqual(self.run_with(dies_after_data)[0], "partial")

    def test_cancel_without_data_keeps_post_pending(self):
        import threading
        ev = threading.Event()

        def fetcher(page, post, b, state):
            ev.set()
            raise cm.TransientError("x")
        with mock.patch.dict(cm.FETCHERS, {"x": fetcher}), mock.patch.object(cm, "RETRY_WAIT", (0, 0, 0)):
            status, items, _ = cm.collect_post_comments("x", FakePage(), {"id": 1, "link": ""}, 10, {}, ev)
        self.assertIsNone(status)


class NetworkLoopTests(unittest.TestCase):
    def setUp(self):
        self.pages = []

        @contextmanager
        def fake_playwright():
            yield object()

        @contextmanager
        def fake_open_page(p, network):
            page = FakePage()
            self.pages.append(page)
            yield page, "sessao"

        @contextmanager
        def fake_lock(network, cancel_event=None, on_wait=None):
            yield True

        self.patches = [
            mock.patch.object(cm, "sync_playwright", fake_playwright),
            mock.patch.object(cm, "open_page", fake_open_page),
            mock.patch.object(cm, "network_lock", fake_lock),
            mock.patch.object(cm, "_pause", lambda *a, **k: None),
            mock.patch.object(cm, "RETRY_WAIT", (0, 0, 0)),
        ]
        for p in self.patches:
            p.start()

    def tearDown(self):
        for p in self.patches:
            p.stop()

    def run_loop(self, fetcher, n=5):
        results = []
        posts = [{"id": i, "link": f"https://x.com/a/status/{i}", "comments": 1} for i in range(n)]
        with mock.patch.dict(cm.FETCHERS, {"x": fetcher}):
            cm.collect_network_comments("x", posts, 10, lambda post, s, items, msg: results.append((post["id"], s, len(items))),
                                        lambda *a, **k: None)
        return results

    def test_every_post_gets_exactly_one_result(self):
        def fetcher(page, post, b, state):
            if post["id"] == 2:
                raise cm.PostUnavailable("removido")
            b.add("c", "@a", "ok")
            b.exhausted = True
        results = self.run_loop(fetcher)
        self.assertEqual([r[0] for r in results], [0, 1, 2, 3, 4])
        self.assertEqual(results[2][1], "error")
        self.assertEqual(len(self.pages), 1)  # um navegador para a rede inteira

    def test_login_marks_remaining_without_visiting(self):
        visited = []

        def fetcher(page, post, b, state):
            visited.append(post["id"])
            if post["id"] == 1:
                raise cm.LoginRequired("sessão expirou")
            b.exhausted = True
        results = self.run_loop(fetcher)
        self.assertEqual(visited, [0, 1])
        self.assertEqual([r[1] for r in results], ["empty", "login", "login", "login", "login"])

    def test_browser_crash_is_recovered(self):
        crashed = []

        def fetcher(page, post, b, state):
            if post["id"] == 1 and not crashed:
                crashed.append(1)
                page.closed = True
                raise RuntimeError("browser has disconnected")
            b.add("c", "@a", "ok")
            b.exhausted = True
        results = self.run_loop(fetcher)
        self.assertEqual([r[1] for r in results], ["done"] * 5)
        self.assertEqual(len(self.pages), 2)

    def test_post_that_always_crashes_is_skipped(self):
        def fetcher(page, post, b, state):
            if post["id"] == 0:
                page.closed = True
                raise RuntimeError("Target closed")
            b.exhausted = True
        results = self.run_loop(fetcher, n=3)
        self.assertEqual([r[:2] for r in results], [(0, "error"), (1, "empty"), (2, "empty")])

    def test_circuit_breaker_stops_endless_failures(self):
        calls = []

        def fetcher(page, post, b, state):
            calls.append(post["id"])
            raise cm.TransientError("HTTP 503")
        results = self.run_loop(fetcher, n=30)
        self.assertEqual(len(results), 30)
        self.assertTrue(all(r[1] == "error" for r in results))
        # 3 navegadores × 4 falhas seguidas × 3 tentativas — o resto é marcado sem abrir
        self.assertEqual(len(calls), (cm.MAX_BROWSER_RESTARTS + 1) * cm.MAX_CONSECUTIVE_FAILURES * cm.MAX_ATTEMPTS)


class PersistenceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        cls.engine = create_engine(f"sqlite:///{os.path.join(cls.tmp, 'test.db')}")
        database.Base.metadata.create_all(cls.engine)
        database.SessionLocal.configure(bind=cls.engine)

    @classmethod
    def tearDownClass(cls):
        database.SessionLocal.configure(bind=database.engine)
        cls.engine.dispose()

    def setUp(self):
        with database.session_scope() as db:
            db.query(database.PostComment).delete()
            db.query(database.ScrapedPost).delete()
            db.add_all([
                database.ScrapedPost(id=1, run_id="run_t", network="YOUTUBE", link="https://www.youtube.com/watch?v=dQw4w9WgXcQ", comments="0"),
                database.ScrapedPost(id=2, run_id="run_t", network="X (Twitter)", link="https://x.com/a/status/2", comments="0"),
                database.ScrapedPost(id=3, run_id="run_t", network="TIKTOK", link="https://www.tiktok.com/@a/video/3", comments="40"),
                database.ScrapedPost(id=4, run_id="run_t", network="INSTAGRAM", link="https://www.instagram.com/p/DdLGFa1AGOv/", comments="9"),
                database.ScrapedPost(id=5, run_id="run_t", network="FACEBOOK", link="", comments="3"),
            ])

    def test_save_is_idempotent_and_keeps_previous_success(self):
        c = [{"id": "a", "author": "@x", "text": "um", "likes": 2}, {"id": "b", "author": "@y", "text": "dois"}]
        self.assertEqual(database.save_post_comments(3, "run_t", "tiktok", c, "done"), 2)
        self.assertEqual(database.save_post_comments(3, "run_t", "tiktok", c + [{"id": "c", "text": "três"}], "done"), 3)
        # recoleta que falhou não apaga nem rebaixa para "error"
        self.assertEqual(database.save_post_comments(3, "run_t", "tiktok", [], "error", "HTTP 500"), 3)
        with database.session_scope() as db:
            post = db.get(database.ScrapedPost, 3)
            self.assertEqual((post.comments_status, post.comments_collected), ("partial", 3))
            self.assertEqual(db.query(database.PostComment).count(), 3)

    def test_select_posts_and_summary(self):
        import comments_service as svc
        with database.session_scope() as db:
            grouped = svc.select_posts(db, "run_t")
            # X com 0 comentários é pulado; YouTube não (a busca não mede comentários); Facebook sem link fica de fora
            self.assertEqual({k: [p["id"] for p in v] for k, v in grouped.items()}, {"youtube": [1], "tiktok": [3], "instagram": [4]})
            self.assertIn("x", svc.select_posts(db, "run_t", skip_zero=False))
            self.assertEqual(list(svc.select_posts(db, "run_t", networks=["instagram"])), ["instagram"])
            self.assertEqual([p["id"] for p in svc.select_posts(db, "run_t", networks=["instagram"])["instagram"]], [4])
            self.assertEqual(list(svc.select_posts(db, "run_t", max_posts=1)), ["tiktok"])  # mais comentado primeiro
        database.save_post_comments(3, "run_t", "tiktok", [{"id": "a", "text": "oi"}], "done")
        database.save_post_comments(1, "run_t", "youtube", [], "disabled")
        with database.session_scope() as db:
            self.assertEqual(list(svc.select_posts(db, "run_t")), ["instagram"])  # concluídos saem dos pendentes
            self.assertIn("tiktok", svc.select_posts(db, "run_t", mode="all"))
            self.assertIn("tiktok", svc.select_posts(db, "run_t", post_ids=[3]))
            s = svc.summary(db, "run_t")
            self.assertEqual(s["totals"]["comments"], 1)
            self.assertEqual(s["totals"]["pending"], 1)
            self.assertEqual(s["totals"]["pending_zero"], 1)
            self.assertEqual(s["statuses"][3][:2], ["done", 1])
            listing = svc.list_comments(db, "run_t", search="oi")
            self.assertEqual(listing["total"], 1)
            self.assertEqual(listing["items"][0]["post"]["id"], 3)
            rows = svc.excel_rows(db, "run_t")
        path = os.path.join(self.tmp, "c.xlsx")
        from excel_builder import build_executive_workbook
        build_executive_workbook([{"Link": "https://x.com/a/status/2", "Rede": "X", "Texto": "t"}], path, "q", rows + [
            {"post_link": "javascript:x", "network": "X", "author": "@b", "text": "=CMD()", "date": "há 2 dias", "is_reply": True}])
        ws = load_workbook(path)["Comentários"]
        self.assertEqual(ws["F2"].value, "oi")
        self.assertEqual(ws["F3"].data_type, "s")  # fórmula injetada vira texto
        self.assertEqual(ws["I3"].value, "há 2 dias")
        self.assertEqual(ws["J3"].value, "Resposta")

    def test_select_posts_balanced_across_networks(self):
        import comments_service as svc
        with database.session_scope() as db:
            db.add_all([
                database.ScrapedPost(id=10, run_id="run_bal", network="INSTAGRAM", link="https://instagram.com/p/1", comments="500"),
                database.ScrapedPost(id=11, run_id="run_bal", network="INSTAGRAM", link="https://instagram.com/p/2", comments="400"),
                database.ScrapedPost(id=12, run_id="run_bal", network="INSTAGRAM", link="https://instagram.com/p/3", comments="300"),
                database.ScrapedPost(id=13, run_id="run_bal", network="INSTAGRAM", link="https://instagram.com/p/4", comments="200"),
                database.ScrapedPost(id=20, run_id="run_bal", network="X (Twitter)", link="https://x.com/u/status/1", comments="15"),
                database.ScrapedPost(id=21, run_id="run_bal", network="X (Twitter)", link="https://x.com/u/status/2", comments="10"),
                database.ScrapedPost(id=22, run_id="run_bal", network="X (Twitter)", link="https://x.com/u/status/3", comments="5"),
            ])
        with database.session_scope() as db:
            # Sem rede especificada e com max_posts=4:
            # Seleciona 2 do Instagram e 2 do X de forma equilibrada, sem que Instagram monopolize
            res = svc.select_posts(db, "run_bal", max_posts=4)
            self.assertIn("instagram", res)
            self.assertIn("x", res)
            self.assertEqual(len(res["instagram"]), 2)
            self.assertEqual(len(res["x"]), 2)
            self.assertEqual([p["id"] for p in res["instagram"]], [10, 11])
            self.assertEqual([p["id"] for p in res["x"]], [20, 21])

            # Quando rede é selecionada diretamente, aplica o limite a ela
            res_x = svc.select_posts(db, "run_bal", networks=["x"], max_posts=2)
            self.assertEqual(list(res_x.keys()), ["x"])
            self.assertEqual([p["id"] for p in res_x["x"]], [20, 21])

            # Modo 'all' também seleciona equilibrado respeitando max_posts mesmo com posts concluídos
            db.get(database.ScrapedPost, 10).comments_status = "done"
            db.get(database.ScrapedPost, 20).comments_status = "empty"
            db.commit()
            res_all = svc.select_posts(db, "run_bal", mode="all", max_posts=4)
            self.assertEqual([p["id"] for p in res_all["instagram"]], [10, 11])
            self.assertEqual([p["id"] for p in res_all["x"]], [20, 21])

    def test_select_posts_recollection_and_network_normalization(self):
        import comments_service as svc
        with database.session_scope() as db:
            db.add_all([
                # Post do X que tinha 0 na métrica inicial mas coletou 25 comentários
                database.ScrapedPost(id=30, run_id="run_norm", network="X (Twitter)", link="https://x.com/u/status/30",
                                     comments="0", comments_collected=25, likes="10", comments_status="done"),
                database.ScrapedPost(id=31, run_id="run_norm", network="X (Twitter)", link="https://x.com/u/status/31",
                                     comments="5", comments_collected=0, likes="2"),
                database.ScrapedPost(id=40, run_id="run_norm", network="INSTAGRAM", link="https://instagram.com/p/40",
                                     comments="100", likes="50"),
                database.ScrapedPost(id=41, run_id="run_norm", network="INSTAGRAM", link="https://instagram.com/p/41",
                                     comments="50", likes="20"),
            ])
        with database.session_scope() as db:
            # 1. Normalização de rede com nome de exibição ("X (Twitter)")
            res_x = svc.select_posts(db, "run_norm", networks=["X (Twitter)"])
            self.assertIn("x", res_x)
            self.assertEqual([p["id"] for p in res_x["x"]], [31])  # 30 está done, modo padrão pending

            # 2. Modo 'all': post com 0 na busca mas 25 coletados não é pulado por skip_zero e é ordenado como 25
            res_all = svc.select_posts(db, "run_norm", networks=["x"], mode="all", skip_zero=True)
            self.assertEqual([p["id"] for p in res_all["x"]], [30, 31])
            self.assertEqual(res_all["x"][0]["comments"], 25)

            # 3. Distribuição assimétrica: Instagram tem 2 posts, X tem 2 posts, max_posts=3
            # Seleciona Instagram (top 1) e X (top 1), e a 3ª vaga vai para o próximo maior
            res_bal = svc.select_posts(db, "run_norm", mode="all", max_posts=3)
            total = sum(len(v) for v in res_bal.values())
            self.assertEqual(total, 3)
            self.assertIn("instagram", res_bal)
            self.assertIn("x", res_bal)
            self.assertIn(30, [p["id"] for p in res_bal["x"]])


if __name__ == "__main__":
    unittest.main()
