"""
Testes unitários e de integração para os novos recursos multimídia,
sanitização de logos, transcrição local e métricas EBM Quintto (Zero Token).
"""
import os
import tempfile
import unittest
from openpyxl import load_workbook
from fastapi.testclient import TestClient

from analytics import build_dashboard, normalize_post, safe_load_media_urls
from brand_logo_service import sanitizar_svg, salvar_logo_marca
from database import Base, Brand, ScrapedPost, ScrapeRun, engine, session_scope
from excel_builder import build_executive_workbook
from main import app
from transcription_service import ALUCINACOES, tem_audio, transcrever_audio_local
from utils import parse_metricas_og_description, sanitizar_metricas_ebm
from video_service import deduplicar_linhas_ocr, duracao_video


class BrandLogoSanitizationTests(unittest.TestCase):
    def test_reject_script_in_svg(self):
        malicious = '<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>'
        valido, erro, cor = sanitizar_svg(malicious)
        self.assertFalse(valido)
        self.assertIn("perigoso", erro.lower())

    def test_reject_foreign_object(self):
        malicious = '<svg xmlns="http://www.w3.org/2000/svg"><foreignObject><div>xss</div></foreignObject></svg>'
        valido, erro, cor = sanitizar_svg(malicious)
        self.assertFalse(valido)
        self.assertIn("perigoso", erro.lower())

    def test_reject_javascript_uri(self):
        malicious = '<svg xmlns="http://www.w3.org/2000/svg"><a href="javascript:alert(1)"><circle r="5"/></a></svg>'
        valido, erro, cor = sanitizar_svg(malicious)
        self.assertFalse(valido)

    def test_reject_external_http_resource(self):
        malicious = '<svg xmlns="http://www.w3.org/2000/svg"><image href="http://evil.com/x.png"/></svg>'
        valido, erro, cor = sanitizar_svg(malicious)
        self.assertFalse(valido)

    def test_reject_invalid_xml(self):
        broken = '<svg><path d="M0 0"</svg>'
        valido, erro, cor = sanitizar_svg(broken)
        self.assertFalse(valido)
        self.assertIn("xml", erro.lower())

    def test_accept_clean_svg(self):
        clean = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100"><circle cx="50" cy="50" r="40" fill="#0D9488"/></svg>'
        valido, resultado, cor = sanitizar_svg(clean)
        self.assertTrue(valido)
        self.assertEqual(cor, "#0D9488")

    def test_reject_doctype_and_entity(self):
        dtd = '<!DOCTYPE svg [<!ENTITY xxe "test">]><svg><text>&xxe;</text></svg>'
        valido, erro, _ = sanitizar_svg(dtd)
        self.assertFalse(valido)
        self.assertIn("dtd/entity", erro.lower())

    def test_reject_xlink_href_external(self):
        malicious = "<svg xmlns:xlink='http://www.w3.org/1999/xlink'><a xlink:href='https://evil.com/x'><circle r='5'/></a></svg>"
        valido, erro, _ = sanitizar_svg(malicious)
        self.assertFalse(valido)
        self.assertTrue("perigoso" in erro.lower() or "externo" in erro.lower())

    def test_salvar_logo_marca_validates_hex(self):
        clean = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100"><rect width="10" height="10" fill="#123456"/></svg>'
        # Cor inválida
        with self.assertRaises(ValueError):
            salvar_logo_marca("MarcaTesteErro", clean, "red")

        # Cor válida
        res = salvar_logo_marca("MarcaTesteOk", clean, "#1E293B")
        self.assertEqual(res["marca"], "MarcaTesteOk")
        self.assertEqual(res["cor"], "#1E293B")


class TranscriptionServiceTests(unittest.TestCase):
    def test_tem_audio_nonexistent(self):
        self.assertFalse(tem_audio("arquivo_que_nao_existe_123.mp4"))

    def test_transcrever_nonexistent(self):
        res = transcrever_audio_local("arquivo_que_nao_existe_123.mp4")
        self.assertEqual(res["status"], "failed")

    def test_anti_hallucination_regex(self):
        self.assertTrue(ALUCINACOES.search("Muito obrigado por assistir!"))
        self.assertTrue(ALUCINACOES.search("Inscreva-se no canal"))
        self.assertTrue(ALUCINACOES.search("deixe seu like"))
        self.assertTrue(ALUCINACOES.search("subtitles by amara.org"))
        self.assertFalse(ALUCINACOES.search("A reunião do comitê municipal tratou de saúde pública."))


class VideoServiceTests(unittest.TestCase):
    def test_duracao_nonexistent(self):
        self.assertEqual(duracao_video("nao_existe.mp4"), 0.0)

    def test_deduplicar_linhas_ocr(self):
        quadros = [
            {"arquivo": "q1.jpg", "texto": ["PREFEITURA DE FORTALEZA"]},
            {"arquivo": "q2.jpg", "texto": ["PREFEITURA DE FORTALEZA"]},
            {"arquivo": "q3.jpg", "texto": ["NOVA OBRA NA BEIRA MAR"]},
        ]
        res = deduplicar_linhas_ocr(quadros)
        self.assertEqual(res["total_linhas_unicas"], 2)
        self.assertIn("PREFEITURA DE FORTALEZA", res["texto_consolidado"])
        self.assertIn("NOVA OBRA NA BEIRA MAR", res["texto_consolidado"])


class MetricIntelligenceTests(unittest.TestCase):
    def test_parse_og_description(self):
        og1 = "1,234 Likes, 56 Comments - Perfil no Instagram: 'Texto do post'"
        parsed1 = parse_metricas_og_description(og1)
        self.assertIsNotNone(parsed1)
        l1, c1 = parsed1
        self.assertEqual(l1, 1234)
        self.assertEqual(c1, 56)

        og2 = "89 curtidas, 12 comentários - Fulano"
        parsed2 = parse_metricas_og_description(og2)
        self.assertIsNotNone(parsed2)
        l2, c2 = parsed2
        self.assertEqual(l2, 89)
        self.assertEqual(c2, 12)

    def test_sanitizar_metricas_corrupted_filler(self):
        # Filler corrompido (views=1000, likes=50)
        res = sanitizar_metricas_ebm(views=1000, likes=50, comments=0, shares=0)
        self.assertEqual(res["status"], "ausente")
        self.assertIsNone(res["views"])
        self.assertIsNone(res["likes"])

    def test_sanitizar_metricas_rejects_synthetic_views(self):
        # Dominação IA: estimativas sintéticas (views = likes * 15) são rejeitadas
        res = sanitizar_metricas_ebm(views=3000, likes=200, comments=10, shares=5)
        self.assertEqual(res["status"], "coletado")
        self.assertEqual(res["likes"], 200)
        self.assertEqual(res["comments"], 10)
        self.assertIsNone(res["views"])

    def test_normalize_post_sanitizes_metrics(self):
        post = ScrapedPost(
            id=999, run_id="r1", network="INSTAGRAM",
            likes="50", views="1000", comments="0", shares="0"
        )
        norm = normalize_post(post)
        self.assertEqual(norm["metrics_status"], "ausente")
        self.assertIsNone(norm["views"])
        self.assertEqual(norm["views_str"], "-")

    def test_build_dashboard_handles_none_views_without_type_error(self):
        with session_scope() as db:
            post = ScrapedPost(
                id=9999, run_id="run_test_views_none", network="INSTAGRAM",
                author="@teste", text="Post com views ausentes",
                likes="200", views="200", comments="10", shares="2"  # views == likes => views vira None
            )
            db.merge(post)

        with session_scope() as db:
            dash = build_dashboard(db, "run_test_views_none")
            self.assertIsNotNone(dash)
            self.assertEqual(dash["summary"]["total_posts"], 1)
            self.assertEqual(dash["summary"]["total_likes"], 200)
            self.assertEqual(dash["summary"]["total_views"], 0)

    def test_safe_load_media_urls_handles_invalid_json(self):
        res = safe_load_media_urls("C:/invalid/path.jpg")
        self.assertEqual(res, ["/api/media/path.jpg"])
        res_empty = safe_load_media_urls("{bad_json")
        self.assertEqual(res_empty, [])


class ExcelTranscriptionSheetTests(unittest.TestCase):
    def test_workbook_with_transcription_creates_extra_sheet(self):
        records = [
            {
                "Link": "https://www.instagram.com/reel/C123456/",
                "Data": "2026-10-01 12:00:00",
                "Rede": "INSTAGRAM",
                "Termo Específico": "Fortaleza",
                "Palavra-Chave": "fortaleza",
                "Autor": "@prefeituradefortaleza",
                "Texto": "Anúncio oficial de novas creches",
                "Curtidas": 1500,
                "Comentarios": 80,
                "Views": 22500,
                "Compartilhamentos": 45,
                "Status Métricas": "exato",
                "media_type": "video",
                "contact_sheet_path": "C:/fake/contact_1.jpg",
                "transcription_text": "Hoje estamos entregando a décima creche em tempo integral para a população.",
                "transcription_status": "done",
                "accessibility_caption": "Prefeito inaugurando prédio escolar moderno com crianças e professores",
            }
        ]
        path = os.path.join(tempfile.mkdtemp(), "teste_multimedia.xlsx")
        build_executive_workbook(records, path, "Fortaleza")
        wb = load_workbook(path)
        self.assertIn("Transcrições & Mídias", wb.sheetnames)
        ws_t = wb["Transcrições & Mídias"]
        self.assertEqual(ws_t["A1"].value, "Link Original")
        self.assertEqual(ws_t["G1"].value, "Transcrição de Áudio")
        self.assertIn("décima creche", ws_t["G2"].value)
        self.assertEqual(ws_t["H2"].value, "contact_1.jpg")


class ApiEndpointsTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def test_get_media_traversal_blocked(self):
        res = self.client.get("/api/media/..secret.txt")
        self.assertEqual(res.status_code, 400)

    def test_get_media_not_found(self):
        res = self.client.get("/api/media/nao_existe_12345.jpg")
        self.assertEqual(res.status_code, 404)

    def test_brand_logo_api_flow(self):
        # 1. Tenta upload de SVG malicioso
        res_mal = self.client.post("/api/brands/PrefeituraFortaleza/logo", json={
            "svg": "<svg><script>alert(1)</script></svg>",
            "color": "#0284C7"
        })
        self.assertEqual(res_mal.status_code, 400)

        # 2. Upload de SVG limpo
        clean_svg = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100"><circle cx="50" cy="50" r="40" fill="#0284C7"/></svg>'
        res_ok = self.client.post("/api/brands/PrefeituraFortaleza/logo", json={
            "svg": clean_svg,
            "color": "#0284C7"
        })
        self.assertEqual(res_ok.status_code, 200)

        # 3. Consulta de SVG salvo
        res_get = self.client.get("/api/brands/PrefeituraFortaleza/logo")
        self.assertEqual(res_get.status_code, 200)
        data = res_get.json()
        self.assertEqual(data["color"], "#0284C7")
        self.assertIn("<circle", data["svg"])
