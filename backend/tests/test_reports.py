import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("RASP_DOWNLOADS_DIR", tempfile.mkdtemp())
os.environ.setdefault("RASP_LOGS_DIR", tempfile.mkdtemp())

from openpyxl import load_workbook  # noqa: E402

from excel_builder import build_executive_workbook  # noqa: E402
from pdf_generator import generate_pdf_report  # noqa: E402

RECORDS = [
    {"Link": "https://x.com/a/status/1", "Data": "2026-09-08 19:20:46", "Rede": "X (Twitter)", "Termo Específico": "Ciro",
     "Palavra-Chave": "ciro OR elmano", "Autor": "@a", "Texto": "=HYPERLINK(\"http://x\",\"x\") \x07 texto", "Curtidas": "1,2 mil",
     "Comentarios": "3", "Views": "39K", "Compartilhamentos": 5},
    {"Link": "javascript:alert(1)", "Data": "há 2 dias", "Rede": "TIKTOK", "Termo Específico": "Elmano, Ciro",
     "Palavra-Chave": "ciro OR elmano", "Autor": "@b*?", "Texto": "Outro post & <tag>", "Curtidas": 10, "Comentarios": 0,
     "Views": "N/A", "Compartilhamentos": 0},
]


class ExcelTests(unittest.TestCase):
    def test_workbook(self):
        path = os.path.join(tempfile.mkdtemp(), "teste.xlsx")
        build_executive_workbook(RECORDS, path, "ciro OR elmano")
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["Dashboard Executivo", "Share of Voice", "Top Autores", "Base de Posts"])
        ws = wb["Base de Posts"]
        self.assertEqual(ws["G2"].data_type, "s")  # conteúdo iniciado por "=" não vira fórmula
        self.assertNotIn("\x07", ws["G2"].value)
        self.assertEqual(ws["H2"].value, 1200)
        self.assertEqual(ws["J2"].value, 39000)
        self.assertEqual(ws["K2"].value, "=H2+I2")
        self.assertEqual(ws["A3"].value, "-")  # link não-http não vira hyperlink
        self.assertEqual(ws.max_column, 12)  # A até L
        self.assertIsNone(ws["M3"].value)
        self.assertTrue(ws["A2"].value.startswith("=HYPERLINK("))

    def test_empty_workbook(self):
        path = os.path.join(tempfile.mkdtemp(), "vazio.xlsx")
        build_executive_workbook([], path, "nada")
        self.assertTrue(os.path.getsize(path) > 0)


class NormalizePostTests(unittest.TestCase):
    def _post(self, **kw):
        from database import ScrapedPost
        base = dict(id=1, run_id="run_1", network="INSTAGRAM", keyword="debate", specific_term="Ciro", author="@a",
                    text="", likes="0", comments="0", views="0", shares="0", link="", post_date="")
        base.update(kw)
        return ScrapedPost(**base)

    def test_numbers_in_text_do_not_become_likes(self):
        from analytics import normalize_post
        item = normalize_post(self._post(author="@jornaldacidadeonline", likes="317", comments="21",
                                         text="Ciro acusa Elmano de repassar R$ 40 milhões a entidades"))
        self.assertEqual(item["likes"], 317)
        self.assertEqual(item["author"], "@jornaldacidadeonline")
        self.assertIn("40 milhões", item["text"])
        zero = normalize_post(self._post(likes="0", text="Post com 5 mil curtidas no texto"))
        self.assertEqual(zero["likes"], 0)

    def test_legacy_missing_likes_recovered_from_text(self):
        from analytics import normalize_post
        item = normalize_post(self._post(likes="N/A", text="perfil • Seguir Áudio original 348 curtidas"))
        self.assertEqual(item["likes"], 348)


class PdfTests(unittest.TestCase):
    def test_pdf_with_markup_characters(self):
        data = {
            "summary": {"total_posts": 2, "total_engagement": 10, "avg_engagement": 5.0},
            "networks": [{"network": "X (TWITTER)", "posts": 2, "likes": 8, "comments": 2, "views": 0, "engagement": 10}],
            "wordcloud": [{"text": "a&b<c>", "weight": 3}],
            "feed": [{"network": "x", "author": "@a<b>", "text": "Texto com & e <script>", "link": "https://x.com/a\"b", "engagement": 10}],
        }
        path = generate_pdf_report(data, "teste_markup.pdf", "consulta & <teste>")
        with open(path, "rb") as f:
            self.assertEqual(f.read(4), b"%PDF")


if __name__ == "__main__":
    unittest.main()
